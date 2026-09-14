#!/usr/bin/env python3
"""Re-evaluate every MUL0912.xlsx row with the current Est_MUL on Ubuntu.

The workbook is updated atomically. Only these result cells are overwritten:

* 自动评估结果: predicted area returned by Est_MUL (um^2)
* 误差: abs(dc综合结果 - prediction) / dc综合结果
* Est_time: per-row Est_MUL wall time (seconds)
* rsq: squared Pearson correlation of all DC and predicted areas (first data row)

Parameter columns, dc综合结果, time, workbook structure, and formatting are kept.

Ubuntu usage (run from this file's directory):

    python3 -m venv .venv-mul0912
    source .venv-mul0912/bin/activate
    pip install -r requirements-mul0912.txt
    python rerun_mul0912.py --backup
"""

from __future__ import annotations

import argparse
import json
import math
import os
import shutil
import sys
import time
from pathlib import Path
from typing import Any

import joblib
import numpy as np
import pandas as pd
from openpyxl import load_workbook

from Est import Est_MUL


ROOT = Path(__file__).resolve().parent
DEFAULT_WORKBOOK = ROOT / "MUL0912.xlsx"
MODEL_ROOT = ROOT / "model"

PARAMETER_HEADERS = (
    "dwt_in_1",
    "dwt_in_2",
    "n_pipelines",
    "sign_in",
    "frac_in_1",
    "frac_in_2",
    "dwt_out",
    "frac_out",
)
RESULT_HEADERS = ("dc综合结果", "自动评估结果", "误差", "time", "Est_time", "rsq")


def _integer(value: Any, header: str, row: int) -> int:
    if isinstance(value, bool):
        return int(value)
    try:
        number = int(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"row {row}: {header} must be an integer, got {value!r}") from error
    if number != value:
        raise ValueError(f"row {row}: {header} must be an integer, got {value!r}")
    return number


def _positive_float(value: Any, header: str, row: int) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"row {row}: {header} must be numeric, got {value!r}") from error
    if not math.isfinite(number) or number <= 0:
        raise ValueError(f"row {row}: {header} must be finite and positive, got {value!r}")
    return number


def _prediction_scalar(value: Any, row: int) -> float:
    values = np.asarray(value, dtype=float).reshape(-1)
    if values.size != 1 or not math.isfinite(float(values[0])) or values[0] <= 0:
        raise ValueError(f"row {row}: Est_MUL returned invalid prediction {value!r}")
    return float(values[0])


def _header_columns(sheet: Any) -> dict[str, int]:
    columns = {
        str(cell.value).strip(): cell.column
        for cell in sheet[1]
        if cell.value is not None and str(cell.value).strip()
    }
    missing = [name for name in (*PARAMETER_HEADERS, *RESULT_HEADERS) if name not in columns]
    if missing:
        raise KeyError(f"worksheet {sheet.title!r} is missing headers: {', '.join(missing)}")
    return columns


def _rsq(actual: list[float], predicted: list[float]) -> float:
    if len(actual) < 2:
        raise ValueError("at least two evaluated rows are required to calculate rsq")
    correlation = float(np.corrcoef(actual, predicted)[0, 1])
    if not math.isfinite(correlation):
        raise ValueError("rsq is undefined because an evaluated series has zero variance")
    return correlation * correlation


def evaluate_workbook(workbook_path: Path, sheet_name: str) -> tuple[Any, dict[str, Any]]:
    model_mul = joblib.load(MODEL_ROOT / "pure_MUL_area.pkl")
    model_su_out = joblib.load(MODEL_ROOT / "SU_out_FxP_area.pkl")
    su_in_db = pd.read_excel(MODEL_ROOT / "SU_in.xlsx", sheet_name="SU_in", header=0)

    workbook = load_workbook(workbook_path, data_only=False)
    if sheet_name not in workbook.sheetnames:
        raise KeyError(f"worksheet {sheet_name!r} not found; available: {workbook.sheetnames}")
    sheet = workbook[sheet_name]
    columns = _header_columns(sheet)

    actual_values: list[float] = []
    predicted_values: list[float] = []
    evaluated_rows: list[int] = []
    total_started = time.perf_counter()

    for row in range(2, sheet.max_row + 1):
        raw_parameters = [sheet.cell(row, columns[name]).value for name in PARAMETER_HEADERS]
        if all(value is None for value in raw_parameters):
            continue
        if any(value is None for value in raw_parameters):
            missing = [
                name for name, value in zip(PARAMETER_HEADERS, raw_parameters) if value is None
            ]
            raise ValueError(f"row {row}: missing parameters: {', '.join(missing)}")

        params = {
            name: _integer(value, name, row)
            for name, value in zip(PARAMETER_HEADERS, raw_parameters)
        }
        sign = params["sign_in"]
        if sign not in (0, 1):
            raise ValueError(f"row {row}: sign_in must be 0 or 1, got {sign!r}")
        actual = _positive_float(
            sheet.cell(row, columns["dc综合结果"]).value, "dc综合结果", row
        )

        started = time.perf_counter()
        prediction = _prediction_scalar(
            Est_MUL(
                model_mul,
                model_su_out,
                su_in_db,
                params["dwt_in_1"],
                params["frac_in_1"],
                sign,
                params["dwt_in_2"],
                params["frac_in_2"],
                sign,
                params["dwt_out"],
                params["frac_out"],
                params["n_pipelines"],
            ),
            row,
        )
        elapsed_seconds = time.perf_counter() - started

        sheet.cell(row, columns["自动评估结果"]).value = prediction
        sheet.cell(row, columns["误差"]).value = abs(actual - prediction) / actual
        sheet.cell(row, columns["Est_time"]).value = elapsed_seconds
        actual_values.append(actual)
        predicted_values.append(prediction)
        evaluated_rows.append(row)

        if len(evaluated_rows) % 50 == 0:
            print(f"evaluated {len(evaluated_rows)} rows", file=sys.stderr, flush=True)

    if not evaluated_rows:
        raise ValueError("no parameter rows were found")

    rsq = _rsq(actual_values, predicted_values)
    rsq_column = columns["rsq"]
    sheet.cell(evaluated_rows[0], rsq_column).value = rsq
    for row in evaluated_rows[1:]:
        sheet.cell(row, rsq_column).value = None

    summary = {
        "workbook": str(workbook_path),
        "sheet": sheet_name,
        "evaluated_rows": len(evaluated_rows),
        "first_row": evaluated_rows[0],
        "last_row": evaluated_rows[-1],
        "prediction_min_um2": min(predicted_values),
        "prediction_max_um2": max(predicted_values),
        "mean_relative_error": float(np.mean([
            abs(actual - predicted) / actual
            for actual, predicted in zip(actual_values, predicted_values)
        ])),
        "rsq": rsq,
        "total_wall_time_seconds": time.perf_counter() - total_started,
    }
    return workbook, summary


def _save_atomically(workbook: Any, source: Path, destination: Path, backup: bool) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination == source:
        temporary = destination.with_name(f".{destination.stem}.tmp{destination.suffix}")
        try:
            workbook.save(temporary)
            if backup:
                shutil.copy2(source, source.with_suffix(source.suffix + ".bak"))
            os.replace(temporary, destination)
        finally:
            temporary.unlink(missing_ok=True)
    else:
        workbook.save(destination)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Re-evaluate all MUL0912.xlsx parameter rows with the current Est_MUL"
    )
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--sheet", default="MUL")
    parser.add_argument(
        "--output",
        type=Path,
        help="write another workbook; omit to atomically overwrite --workbook",
    )
    parser.add_argument(
        "--backup",
        action="store_true",
        help="when overwriting in place, save the original as MUL0912.xlsx.bak",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    source = args.workbook.expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(f"workbook not found: {source}")
    destination = args.output.expanduser().resolve() if args.output else source
    workbook, summary = evaluate_workbook(source, args.sheet)
    _save_atomically(workbook, source, destination, args.backup)
    summary["output"] = str(destination)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
